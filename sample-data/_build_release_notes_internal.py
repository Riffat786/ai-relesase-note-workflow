# -*- coding: utf-8 -*-
"""Build publication-ready HTML and DOCX release notes for Release 14.1 — INTERNAL edition.
Content is the approved internal text from release-notes-internal.md.
Internal audience: Jira keys, API/DB changes, testing, deployment notes, and open
data-quality actions are intentionally retained (unlike the customer edition).
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT_DIR = os.path.dirname(os.path.abspath(__file__))

TITLE = "Release Notes 14.1 — Internal"

META = [
    ("Release Version", "14.1"),
    ("Release Date", "2026-07-05"),
    ("Audience", "Internal (Engineering, QA, Release Management, Support)"),
    ("Project", "KAN — Release Notes Automation"),
    ("Issues in Release", "KAN-1, KAN-2, KAN-7"),
]

OVERVIEW = ("Release 14.1 delivers one new feature (bulk customer import), one defect fix "
            "(duplicate transaction prevention), and one regulatory enhancement (updated report "
            "templates). All three issues are in Done status. This edition is for internal "
            "distribution and includes implementation, API, database, testing, and deployment "
            "details, plus open data-quality actions.")

# Each item carries full internal detail.
SECTIONS = [
    ("New Features", [
        {
            "title": "KAN-1 — Bulk Customer Import",
            "product": "Business Loan",
            "module": "Customer Details",
            "work_type": "Feature",
            "summary": "CSV-based bulk customer import to reduce manual onboarding effort.",
            "tech_changes": [
                "CSV parsing service.",
                "Customer validation framework.",
                "Duplicate customer detection.",
                "Customer Service API integration.",
                "Upload progress tracking and audit logging.",
            ],
            "api": "POST /customers/import",
            "db": "Added BulkUploadId and UploadStatus columns to the CustomerUpload table.",
            "validation": "Mandatory fields, duplicate customers, invalid CSV format.",
            "testing": "Unit, Integration, and Regression completed/passed.",
            "deployment": "No manual configuration required.",
        },
    ]),
    ("Bug Fixes", [
        {
            "title": "KAN-2 — Duplicate Transaction Validation",
            "product": "Business Loan",
            "module": "Loan Processing",
            "work_type": "Bug",
            "summary": ("Prevents the same loan transaction from being processed more than once by "
                        "validating transaction references before submission."),
            "tech_changes": [
                "Duplicate transaction validation.",
                "Transaction ID comparison before processing.",
                "Duplicate transaction logging.",
                "User-friendly validation message on detection.",
            ],
            "api": "None",
            "db": "Added TransactionHash index.",
            "validation": "Duplicate Transaction ID, duplicate request payload, duplicate reference number.",
            "testing": "Unit, Integration, and Regression completed.",
            "deployment": "⚠ Database index required — must be applied as part of deployment.",
        },
    ]),
    ("Regulatory Updates", [
        {
            "title": "KAN-7 — Update Regulatory Report Template",
            "product": "Not set in Jira",
            "module": "Reports",
            "work_type": "Enhancement",
            "summary": "Regulatory report templates updated to meet revised compliance requirements.",
            "tech_changes": [
                "Updated report layout.",
                "Added new compliance fields.",
                "Improved export performance.",
            ],
            "api": "GET /reports/regulatory",
            "db": "Added ComplianceCategory column.",
            "validation": "Mandatory compliance fields.",
            "testing": "Completed.",
            "deployment": "⚠ Report template deployment required.",
        },
    ]),
    ("Known Issues", None),
]

ACTIONS = [
    ("KAN-7 — Missing Product in Jira", "The Product field (customfield_10074) is not set. Populate before the next extraction."),
    ("KAN-7 — Release/Product discrepancy", ("Jira records Release 14.1 / Product unset, but the implementation note records "
                                             "Release 14.2 / Product Mortgage Loan. Reconcile before publishing customer-facing "
                                             "notes — if KAN-7 belongs to 14.2, remove it from this release.")),
    ("Deployment ordering", ("KAN-2 (DB index) and KAN-7 (report template deployment) both require manual deployment steps — "
                             "confirm these are in the release runbook.")),
]

BULLET = "■"  # solid black square

# ---------------------------------------------------------------- HTML
def esc(s):
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

def build_html():
    css = """
    body { font-family: Cambria, Georgia, serif; color:#000; font-size:11pt; line-height:1.5; margin:40px; }
    h1 { font-family: Cambria, Georgia, serif; font-size:20pt; font-weight:bold; color:#000; text-align:left; }
    h2 { font-family: Cambria, Georgia, serif; font-size:16pt; font-weight:bold; color:#000; text-align:left; margin-top:28px; }
    h3 { font-family: Cambria, Georgia, serif; font-size:12pt; font-weight:bold; color:#000; margin-bottom:4px; }
    table.meta { border-collapse:collapse; margin:16px 0 8px 0; }
    table.meta td { border:1px solid #333; padding:6px 14px; font-size:11pt; }
    table.meta td.k { font-weight:bold; background:#f2f2f2; }
    ul { list-style:none; padding-left:18px; }
    ul li { position:relative; padding-left:20px; margin:4px 0; }
    ul li:before { content:"\\25A0"; position:absolute; left:0; color:#000; }
    .meta-line { margin:2px 0; }
    """
    p = []
    p.append("<!DOCTYPE html>")
    p.append('<html lang="en"><head><meta charset="utf-8">')
    p.append(f"<title>{esc(TITLE)}</title>")
    p.append(f"<style>{css}</style></head><body>")
    p.append(f"<h1>{esc(TITLE)}</h1>")
    p.append('<table class="meta">')
    for k, v in META:
        p.append(f'<tr><td class="k">{esc(k)}</td><td>{esc(v)}</td></tr>')
    p.append("</table>")
    p.append("<h2>Overview</h2>")
    p.append(f"<p>{esc(OVERVIEW)}</p>")
    for name, items in SECTIONS:
        p.append(f"<h2>{esc(name)}</h2>")
        if items is None:
            p.append("<p>None reported for this release.</p>")
            continue
        for it in items:
            p.append(f"<h3>{esc(it['title'])}</h3>")
            p.append(f'<p class="meta-line"><strong>Product / Module:</strong> {esc(it["product"])} / {esc(it["module"])}</p>')
            p.append(f'<p class="meta-line"><strong>Work Type:</strong> {esc(it["work_type"])}</p>')
            p.append(f'<p class="meta-line"><strong>Summary:</strong> {esc(it["summary"])}</p>')
            p.append("<p class=\"meta-line\"><strong>Technical Changes:</strong></p>")
            p.append("<ul>")
            for tc in it["tech_changes"]:
                p.append(f"<li>{esc(tc)}</li>")
            p.append("</ul>")
            p.append("<ul>")
            p.append(f'<li><strong>API Changes:</strong> {esc(it["api"])}</li>')
            p.append(f'<li><strong>Database Changes:</strong> {esc(it["db"])}</li>')
            p.append(f'<li><strong>Validation:</strong> {esc(it["validation"])}</li>')
            p.append(f'<li><strong>Testing:</strong> {esc(it["testing"])}</li>')
            p.append(f'<li><strong>Deployment Notes:</strong> {esc(it["deployment"])}</li>')
            p.append("</ul>")
    # Open actions
    p.append("<h2>Open Actions / Data-Quality Notes</h2>")
    p.append("<ul>")
    for label, text in ACTIONS:
        p.append(f"<li><strong>{esc(label)}:</strong> {esc(text)}</li>")
    p.append("</ul>")
    p.append("</body></html>")
    html = "\n".join(p)
    path = os.path.join(OUT_DIR, "release-notes-internal.html")
    with open(path, "w", encoding="utf-8") as f:
        f.write(html)
    return path

# ---------------------------------------------------------------- DOCX
def set_font(run, size, bold=False):
    run.font.name = "Cambria"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)

def add_para(doc, text, size=11, bold=False, space_after=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    set_font(r, size, bold)
    return p

def add_label_line(doc, label, text, space_after=2):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    rb = p.add_run(label + ": ")
    set_font(rb, 11, True)
    rv = p.add_run(text)
    set_font(rv, 11, False)
    return p

def add_bullet(doc, label, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Pt(18)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(BULLET + "  ")
    set_font(r, 11, False)
    if label:
        rb = p.add_run(label + ": ")
        set_font(rb, 11, True)
    rt = p.add_run(text)
    set_font(rt, 11, False)
    return p

def build_docx(filename="release-notes-internal.docx"):
    doc = Document()
    add_para(doc, TITLE, size=20, bold=True, space_after=12)
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for k, v in META:
        cells = table.add_row().cells
        rk = cells[0].paragraphs[0].add_run(k)
        set_font(rk, 11, True)
        rv = cells[1].paragraphs[0].add_run(v)
        set_font(rv, 11, False)
    doc.add_paragraph()
    add_para(doc, "Overview", size=16, bold=True, space_after=6)
    add_para(doc, OVERVIEW, size=11, bold=False, space_after=10)
    for name, items in SECTIONS:
        add_para(doc, name, size=16, bold=True, space_after=6)
        if items is None:
            add_para(doc, "None reported for this release.", size=11, space_after=10)
            continue
        for it in items:
            add_para(doc, it["title"], size=12, bold=True, space_after=2)
            add_label_line(doc, "Product / Module", f'{it["product"]} / {it["module"]}')
            add_label_line(doc, "Work Type", it["work_type"])
            add_label_line(doc, "Summary", it["summary"])
            add_label_line(doc, "Technical Changes", "")
            for tc in it["tech_changes"]:
                add_bullet(doc, None, tc)
            add_bullet(doc, "API Changes", it["api"])
            add_bullet(doc, "Database Changes", it["db"])
            add_bullet(doc, "Validation", it["validation"])
            add_bullet(doc, "Testing", it["testing"])
            add_bullet(doc, "Deployment Notes", it["deployment"])
            doc.add_paragraph()
    add_para(doc, "Open Actions / Data-Quality Notes", size=16, bold=True, space_after=6)
    for label, text in ACTIONS:
        add_bullet(doc, label, text)
    path = os.path.join(OUT_DIR, filename)
    try:
        doc.save(path)
        return path
    except PermissionError:
        alt = os.path.join(OUT_DIR, "release-notes-internal-published.docx")
        doc.save(alt)
        return alt + "  (original was locked/open in Word)"

if __name__ == "__main__":
    h = build_html()
    d = build_docx()
    print("HTML:", h)
    print("DOCX:", d)
