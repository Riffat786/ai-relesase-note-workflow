# -*- coding: utf-8 -*-
"""Build publication-ready HTML and DOCX release notes for Release 14.1.
Content is the approved customer-facing text from release-notes.md.
No Jira IDs, review artifacts, or implementation details are included.
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT_DIR = os.path.dirname(os.path.abspath(__file__))

TITLE = "Release 14.1 — Customer Onboarding & Compliance Update"

META = [
    ("Release Version", "14.1"),
    ("Release Date", "2026-07-05"),
    ("Audience", "Customer"),
]

OVERVIEW = ("Release 14.1 streamlines customer onboarding, strengthens the reliability of "
            "transaction processing, and keeps regulatory reporting aligned with the latest "
            "compliance requirements. Together, these changes help teams onboard customers "
            "faster, process payments with greater confidence, and stay compliant with "
            "evolving reporting standards.")

# Each item: title, product (or None), module (or None), description, benefit_label, benefit_text
SECTIONS = [
    ("New Features", [
        {
            "title": "Bulk Customer Import",
            "product": "Business Loan",
            "module": "Customer Details",
            "description": ("Onboard multiple customers at once by uploading a single file, instead "
                            "of entering records one at a time. Records are validated as they are "
                            "processed, potential duplicates are flagged, and upload progress is "
                            "tracked throughout."),
            "benefit_label": "Customer Benefit",
            "benefit_text": ("Significantly reduces manual data entry and shortens onboarding time "
                             "while helping ensure clean, accurate customer records."),
        },
    ]),
    ("Bug Fixes", [
        {
            "title": "Duplicate Transaction Prevention",
            "product": "Business Loan",
            "module": "Loan Processing",
            "description": ("Resolved an issue that could allow the same transaction to be processed "
                            "more than once. Each transaction is now verified before submission, and "
                            "a clear message is shown when a potential duplicate is detected."),
            "benefit_label": "Customer Benefit",
            "benefit_text": ("Prevents unintended duplicate payments and improves the accuracy and "
                             "reliability of transaction processing."),
        },
    ]),
    ("Regulatory Updates", [
        {
            "title": "Updated Regulatory Report Templates",
            "product": None,
            "module": "Reports",
            "description": ("Regulatory report templates have been updated to comply with the latest "
                            "reporting standards, including new compliance information and a refreshed "
                            "report layout. Report generation has also been made faster."),
            "benefit_label": "Business Impact",
            "benefit_text": ("Helps the organization meet current regulatory requirements and produce "
                             "compliant reports more efficiently."),
        },
    ]),
    ("Known Issues", None),  # None -> single line "None reported for this release."
]

BULLET = "■"  # solid black square

# ---------------------------------------------------------------- HTML
def build_html():
    css = """
    body { font-family: Cambria, Georgia, serif; color:#000; font-size:11pt; line-height:1.5; margin:40px; }
    h1 { font-family: Cambria, Georgia, serif; font-size:20pt; font-weight:bold; color:#000; text-align:left; }
    h2 { font-family: Cambria, Georgia, serif; font-size:16pt; font-weight:bold; color:#000; text-align:left; margin-top:28px; }
    h3 { font-family: Cambria, Georgia, serif; font-size:12pt; font-weight:bold; color:#000; margin-bottom:4px; }
    table.meta { border-collapse:collapse; margin:16px 0 8px 0; }
    table.meta td { border:1px solid #333; padding:6px 14px; font-size:11pt; }
    table.meta td.k { font-weight:bold; background:#f2f2f2; }
    ul { list-style:none; padding-left:18px; }
    ul li { position:relative; padding-left:20px; margin:4px 0; }
    ul li:before { content:"\\25A0"; position:absolute; left:0; color:#000; }
    .meta-line { margin:2px 0; }
    """
    parts = []
    parts.append("<!DOCTYPE html>")
    parts.append('<html lang="en"><head><meta charset="utf-8">')
    parts.append(f"<title>{TITLE}</title>")
    parts.append(f"<style>{css}</style></head><body>")
    parts.append(f"<h1>{TITLE}</h1>")
    # metadata table
    parts.append('<table class="meta">')
    for k, v in META:
        parts.append(f'<tr><td class="k">{k}</td><td>{v}</td></tr>')
    parts.append("</table>")
    # overview
    parts.append("<h2>Overview</h2>")
    parts.append(f"<p>{OVERVIEW}</p>")
    # sections
    for name, items in SECTIONS:
        parts.append(f"<h2>{name}</h2>")
        if items is None:
            parts.append("<p>None reported for this release.</p>")
            continue
        for it in items:
            parts.append(f"<h3>{it['title']}</h3>")
            if it.get("product"):
                parts.append(f'<p class="meta-line"><strong>Product:</strong> {it["product"]}</p>')
            if it.get("module"):
                parts.append(f'<p class="meta-line"><strong>Module:</strong> {it["module"]}</p>')
            parts.append("<ul>")
            parts.append(f"<li>{it['description']}</li>")
            parts.append(f"<li><strong>{it['benefit_label']}:</strong> {it['benefit_text']}</li>")
            parts.append("</ul>")
    parts.append("</body></html>")
    html = "\n".join(parts)
    path = os.path.join(OUT_DIR, "release-notes.html")
    with open(path, "w", encoding="utf-8") as f:
        f.write(html)
    return path

# ---------------------------------------------------------------- DOCX
def set_font(run, size, bold=False):
    run.font.name = "Cambria"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)

def add_para(doc, text, size=11, bold=False, space_after=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    set_font(r, size, bold)
    return p

def add_bullet(doc, label, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Pt(18)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(BULLET + "  ")
    set_font(r, 11, False)
    if label:
        rb = p.add_run(label + ": ")
        set_font(rb, 11, True)
    rt = p.add_run(text)
    set_font(rt, 11, False)
    return p

def build_docx(filename="release-notes.docx"):
    doc = Document()
    # Title (H1)
    add_para(doc, TITLE, size=20, bold=True, space_after=12)
    # Metadata table
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for k, v in META:
        cells = table.add_row().cells
        rk = cells[0].paragraphs[0].add_run(k)
        set_font(rk, 11, True)
        rv = cells[1].paragraphs[0].add_run(v)
        set_font(rv, 11, False)
    doc.add_paragraph()
    # Overview
    add_para(doc, "Overview", size=16, bold=True, space_after=6)
    add_para(doc, OVERVIEW, size=11, bold=False, space_after=10)
    # Sections
    for name, items in SECTIONS:
        add_para(doc, name, size=16, bold=True, space_after=6)
        if items is None:
            add_para(doc, "None reported for this release.", size=11, space_after=10)
            continue
        for it in items:
            add_para(doc, it["title"], size=12, bold=True, space_after=2)
            if it.get("product"):
                add_bullet_meta = None
                p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
                rb = p.add_run("Product: "); set_font(rb, 11, True)
                rv = p.add_run(it["product"]); set_font(rv, 11, False)
            if it.get("module"):
                p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
                rb = p.add_run("Module: "); set_font(rb, 11, True)
                rv = p.add_run(it["module"]); set_font(rv, 11, False)
            add_bullet(doc, None, it["description"])
            add_bullet(doc, it["benefit_label"], it["benefit_text"])
            doc.add_paragraph()
    # Save (fall back if locked)
    path = os.path.join(OUT_DIR, filename)
    try:
        doc.save(path)
        return path
    except PermissionError:
        alt = os.path.join(OUT_DIR, "release-notes-published.docx")
        doc.save(alt)
        return alt + "  (original was locked/open in Word)"

if __name__ == "__main__":
    h = build_html()
    d = build_docx()
    print("HTML:", h)
    print("DOCX:", d)
